"""Gera a planilha do estoque que o site dos armários lê (para subir no Google Planilhas).

Os móveis vêm de laboratorio.js, na raiz do site. A planilha recebe a lista de móveis para a lista suspensa
e uma aba com os códigos de cada gaveta, prateleira e compartimento.

Uso: python gerar_planilha.py saida.xlsx URL_DO_SITE
"""
import json
import pathlib
import sys
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.datavalidation import DataValidation

saida, site = sys.argv[1], sys.argv[2]

VERDE, CINZA, BORDA = "70AD47", "EFEFEF", "BFBFBF"
F = "Arial"
fonte = Font(name=F, size=10)
cab_fonte = Font(name=F, size=10, bold=True, color="FFFFFF")
cab_fill = PatternFill("solid", fgColor=VERDE)
calc_fill = PatternFill("solid", fgColor=CINZA)
linha = Side(style="thin", color=BORDA)
borda = Border(left=linha, right=linha, top=linha, bottom=linha)
LINHAS = 1000  # linhas preparadas com listas suspensas


def norm(s):
    return unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower().strip()


def ler_laboratorio():
    texto = (pathlib.Path(__file__).resolve().parent.parent / "laboratorio.js").read_text(encoding="utf-8")
    ini = texto.index("{", texto.index("window.LABORATORIO"))
    return json.loads(texto[ini:texto.rindex("}") + 1])


# Mesmos códigos que o site calcula em slotsOf(): P, G e C contados de cima para baixo, uma letra por divisão
PREFIXO = {"prateleira": "P", "gaveta": "G", "compartimento": "C"}
NOME = {"prateleira": "Prateleira ", "gaveta": "Gaveta ", "compartimento": "Compartimento "}
LADOS = {2: ["esquerda", "direita"], 3: ["esquerda", "meio", "direita"]}


def tipo_nivel(s):
    t = norm(s)
    return "gaveta" if t.startswith("g") else "compartimento" if t.startswith("c") else "prateleira"


def codigos(movel):
    caixa = any(p in norm(movel.get("tipo")) for p in ("caixa", "maleta"))
    cont = {k: 0 for k in PREFIXO}
    for nivel in movel.get("niveis") or [{"tipo": "Compartimento" if caixa else "Prateleira"}]:
        tipo = tipo_nivel(nivel.get("tipo"))
        cont[tipo] += 1
        n, div = cont[tipo], max(1, min(6, int(round(float(nivel.get("divisoes") or 1)))))
        base = nivel.get("nome") or NOME[tipo] + str(n)
        for c in range(div):
            if div == 1:
                yield PREFIXO[tipo] + str(n), base
            else:
                lado = LADOS[div][c] if div in LADOS and not caixa else "divisão " + "ABCDEF"[c]
                yield PREFIXO[tipo] + str(n) + "ABCDEF"[c], base + ", " + lado


LAB = ler_laboratorio()
MOVEIS = LAB["armarios"]
ARMARIOS = [m["armario"] for m in MOVEIS]
CLASSES = ["Mecânica", "Eletrônica", "Projeto antigo"]
CATEGORIAS = ["ESC", "Motores", "Baterias", "Carregadores", "Conectores", "Fios e cabos",
              "Componentes (capacitores, resistores...)", "Placas e controladores", "Rádio e receptor",
              "Sensores", "Parafusos, porcas e arruelas", "Rolamentos e buchas", "Eixos e acoplamentos",
              "Rodas", "Chapas e perfis", "Filamento", "Peças impressas", "Ferramentas", "Consumíveis", "Outros"]
ESTADOS = ["Novo", "Usado (funcionando)", "Não testado", "Com defeito"]
UNIDADES = ["un", "par", "m", "kit", "rolo", "pacote", "caixa"]

# O que deu para ver na foto da maleta preta em 26/10/2026. Fica marcado para alguém conferir no laboratório.
FOTO = "Cadastrado pela foto de 26/10/2026. Conferir no laboratório."
ITENS_DAS_FOTOS = [
    ["Maleta 12", "C1A", "Eletrônica", "Motores", "Motor brushless", "1 laranja e 3 pretos, com fios", 4, "un", "", "", "", FOTO],
    ["Maleta 12", "C1B", "Eletrônica", "Motores", "Motor brushless", "1 azul pequeno e outros pretos, com fios", 4, "un", "", "", "", FOTO + " Confira a quantidade."],
    ["Maleta 12", "C1C", "Eletrônica", "Rádio e receptor", "Receptor Radiolink", "Com o manual dos receptores de superfície", 1, "un", "", "", "", FOTO],
    ["Maleta 12", "C1C", "Eletrônica", "Rádio e receptor", "Receptor DSMX", "S603", 1, "un", "", "", "", FOTO],
    ["Maleta 12", "C1D", "", "Outros", "Saquinhos com peças pequenas", "", None, "", "", "", "", FOTO + " Veja o que tem dentro."],
    ["Maleta 12", "C2", "Eletrônica", "Motores", "Motor F2004 Ultralight", "Na caixa", 2, "un", "", "", "", FOTO],
    ["Maleta 12", "C2", "Eletrônica", "Motores", "Motor brushless vermelho", "Com conectores bala dourados", 1, "un", "", "", "", FOTO],
    ["Maleta 12", "C2", "Eletrônica", "Motores", "Motor pequeno em suporte de MDF", "Com cabo de 3 fios", 1, "un", "", "", "", FOTO],
    ["Maleta 12", "C2", "Eletrônica", "Baterias", "Pilha AA alcalina", "i-Tech, 2 cartelas de 4", 8, "un", "Novo", "", "", FOTO],
]


def cabecalho(ws, nomes, larguras):
    for i, (nome, larg) in enumerate(zip(nomes, larguras), start=1):
        c = ws.cell(row=1, column=i, value=nome)
        c.font, c.fill, c.border = cab_fonte, cab_fill, borda
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = larg
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def lista(ws, col, faixa, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                        errorTitle="Valor fora da lista", error="Escolha uma opção da lista. As opções ficam na aba Listas.")
    dv.add(f"{col}2:{col}{faixa}")
    ws.add_data_validation(dv)


wb = Workbook()
# Arial 10 como fonte padrão do arquivo, sem precisar formatar célula por célula
wb._fonts = IndexedList([Font(name=F, size=10)])

# ---- Como usar ----
ws = wb.active
ws.title = "Como usar"
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 110
textos = [
    ("Estoque do laboratório NEON", Font(name=F, size=16, bold=True)),
    ("Esta planilha alimenta o modelo 3D dos armários: " + site, Font(name=F, size=10, color="4F8F2C", underline="single")),
    ("", fonte),
    ("Como cadastrar um item", Font(name=F, size=12, bold=True)),
    ("1. Na aba Estoque, escreva uma linha por item. As colunas Armário e Local no armário dizem onde ele está.", fonte),
    ("2. Local no armário é o código do compartimento: G1 é a gaveta 1, P2 é a prateleira 2, P2A é o lado esquerdo da prateleira 2 e C1 é o compartimento 1 de uma caixa ou maleta.", fonte),
    ("    Os códigos de cada móvel estão na aba Locais. No site, clicar num compartimento também mostra o código dele.", fonte),
    ("3. Em Qtd, escreva só o número (10, não \"10 un\"). A unidade vai na coluna ao lado.", fonte),
    ("4. O site lê a planilha sozinho a cada minuto. Para ver na hora, clique em Atualizar no site.", fonte),
    ("", fonte),
    ("Exemplo de linha na aba Estoque", Font(name=F, size=12, bold=True)),
    ("Gaveteiro 6  |  G1  |  Eletrônica  |  Conectores  |  XT60 fêmea  |  Amass, original  |  10  |  un  |  Novo  |  (vazio)  |  26/10/2026  |  (vazio)", fonte),
    ("", fonte),
    ("Os móveis do modelo", Font(name=F, size=12, bold=True)),
    ("Os móveis, as medidas e as gavetas de cada um ficam no próprio site, no arquivo laboratorio.js, e não nesta planilha.", fonte),
    ("Se um móvel mudar de lugar ou ganhar uma prateleira, ajuste esse arquivo e gere esta planilha de novo, para as abas Locais e Listas acompanharem.", fonte),
    ("Os itens marcados como \"Cadastrado pela foto\" foram lidos das fotos do laboratório. Confira e apague a observação.", fonte),
    ("", fonte),
    ("Não mude os nomes das abas nem os títulos das colunas da aba Estoque: o site procura por eles.", fonte),
]
for i, (t, f) in enumerate(textos, start=1):
    c = ws.cell(row=i, column=2, value=t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws["B2"].hyperlink = site

# ---- Estoque ----
ws = wb.create_sheet("Estoque")
cols = ["Armário", "Local no armário", "Classificação", "Categoria", "Item", "Especificação / modelo",
        "Qtd", "Unid.", "Estado", "Projeto de origem", "Conferido em", "Observações"]
cabecalho(ws, cols, [14, 12, 15, 26, 30, 32, 7, 8, 18, 18, 13, 40])
# formato de texto na coluna inteira, para G1 e P2B nunca virarem número ou data
ws.column_dimensions["B"].number_format = "@"
for r, valores in enumerate(ITENS_DAS_FOTOS, start=2):
    for c, v in enumerate(valores, start=1):
        cell = ws.cell(row=r, column=c, value=v if v != "" else None)
        cell.font = fonte
        if c == 2:
            cell.number_format = "@"
lista(ws, "A", LINHAS, "=Listas!$A$2:$A$60")
lista(ws, "C", LINHAS, "=Listas!$B$2:$B$30")
lista(ws, "D", LINHAS, "=Listas!$C$2:$C$40")
lista(ws, "H", LINHAS, "=Listas!$E$2:$E$30")
lista(ws, "I", LINHAS, "=Listas!$D$2:$D$30")
ws.auto_filter.ref = f"A1:L{LINHAS}"

# ---- Locais ----
ws = wb.create_sheet("Locais")
cabecalho(ws, ["Armário", "Nome", "Local no armário", "O que é"], [14, 26, 16, 34])
r = 2
for m in MOVEIS:
    for codigo, desc in codigos(m):
        for c, v in enumerate([m["armario"], m.get("nome", ""), codigo, desc], start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.fill = fonte, calc_fill
        r += 1
ws.auto_filter.ref = f"A1:D{r - 1}"

# ---- Listas ----
ws = wb.create_sheet("Listas")
cabecalho(ws, ["Armários", "Classificação", "Categoria", "Estado", "Unidade"], [14, 16, 36, 20, 10])
for c, valores in enumerate([ARMARIOS, CLASSES, CATEGORIAS, ESTADOS, UNIDADES], start=1):
    for i, v in enumerate(valores, start=2):
        ws.cell(row=i, column=c, value=v).font = fonte

wb.save(saida)

# Tira do pacote as partes opcionais (tema e propriedades) para o arquivo ficar pequeno
import re as _re, zipfile
with zipfile.ZipFile(saida) as z:
    partes = {n: z.read(n) for n in z.namelist()}
for n in list(partes):
    if n.startswith(("xl/theme/", "docProps/")):
        del partes[n]
partes["[Content_Types].xml"] = _re.sub(rb'<Override PartName="/(xl/theme|docProps)/[^>]*/>', b"", partes["[Content_Types].xml"])
partes["_rels/.rels"] = _re.sub(rb'<Relationship [^>]*Target="/?docProps/[^>]*/>', b"", partes["_rels/.rels"])
partes["xl/_rels/workbook.xml.rels"] = _re.sub(rb'<Relationship [^>]*Target="/?(xl/)?theme/[^>]*/>', b"", partes["xl/_rels/workbook.xml.rels"])
with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
    for n, dados in partes.items():
        z.writestr(n, dados)
print("ok", saida, len(MOVEIS), "móveis,", r - 2, "locais,", len(ITENS_DAS_FOTOS), "itens das fotos")
